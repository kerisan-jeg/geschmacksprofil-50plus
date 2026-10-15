"""Lebensmittelliste für die Suche aus der Schweizer Nährwertdatenbank (BLV) erzeugen.

Quelle: Schweizer Nährwertdatenbank, Generische Lebensmittel, Bundesamt für Lebensmittelsicherheit und
Veterinärwesen (BLV), https://naehrwertdaten.ch. Frei verwendbar mit Quellenvermerk.

Aufruf (einmalig, wenn eine neue Version erscheint):
  1. Excel-Datei von https://naehrwertdaten.ch/de/downloads/ herunterladen
  2. pip install openpyxl
  3. python scripts/import-foods.py <Schweizer_Nahrwertdatenbank.xlsx>
Schreibt src/data/foods.json.

Regeln:
  - Zutaten und Getränke; keine Gerichte (Kategorie „Gerichte“, Kuchen und Torten): Gerichte gehören
    nach nurec v0.2 in Stufe 4 und werden dort als eigene Karten bewertet.
  - Varianten werden zusammengefasst: „Ananas, roh“, „Ananas in Sirup, Konserve“ → „Ananas“.
    Weitere Teile mit grossem Anfangsbuchstaben werden Suchbegriffe („Rind, Entrecôte, roh“ → Rind,
    gefunden auch über „Entrecôte“), ebenso die Synonyme der Datenbank.
  - Jede Gruppe bekommt einen kurzen Namen für die Anzeige („Olive · Gemüse“).
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent.parent / 'src' / 'data' / 'foods.json'

# Gerichte und gebackene Desserts gehören in Stufe 4, nicht in die Suche nach Zutaten.
EXCLUDED = ('Gerichte', 'Speziallebensmittel', 'Süssigkeiten/Kuchen, Torten und Cake')

# Anzeige-Gruppe je Kategorie der Datenbank (genaueste zuerst).
GROUPS = [
    ('Gemüse/Kräuter', 'Kräuter'),
    ('Gemüse/Pilze', 'Pilze'),
    ('Gemüse/Gemüsesäfte', 'Getränke'),
    ('Gemüse', 'Gemüse'),
    ('Früchte/Fruchtsäfte', 'Getränke'),
    ('Früchte', 'Obst'),
    ('Fleisch und Innereien', 'Fleisch und Geflügel'),
    ('Fleisch- und Wurstwaren', 'Wurstwaren'),
    ('Fisch', 'Fisch und Meeresfrüchte'),
    ('Milch und Milchprodukte', 'Milchprodukte'),
    ('Fette und Öle/Rahm', 'Milchprodukte'),
    ('Eier', 'Eier'),
    ('Getreideprodukte, Hülsenfrüchte und Kartoffeln/Hülsenfrüchte', 'Hülsenfrüchte'),
    ('Getreideprodukte, Hülsenfrüchte und Kartoffeln/Kartoffeln', 'Kartoffeln'),
    ('Getreideprodukte, Hülsenfrüchte und Kartoffeln/Teigwaren', 'Pasta'),
    ('Getreideprodukte, Hülsenfrüchte und Kartoffeln/Mehle', 'Mehl'),
    ('Getreideprodukte, Hülsenfrüchte und Kartoffeln/Teige', 'Teig'),
    ('Getreideprodukte, Hülsenfrüchte und Kartoffeln', 'Reis und andere Getreide'),
    ('Brote, Flocken und Frühstückscerealien/Flocken', 'Flocken und Müesli'),
    ('Brote, Flocken und Frühstückscerealien/Müesli', 'Flocken und Müesli'),
    ('Brote, Flocken und Frühstückscerealien', 'Brot'),
    ('Nüsse, Samen und Ölfrüchte', 'Nüsse und Samen'),
    ('Fette und Öle', 'Fette, Öle und Saucen'),
    ('Süssigkeiten', 'Süsses'),
    ('Salzige Snacks', 'Snacks'),
    ('Pflanzliche Proteinlieferanten und Alternativen zu tierischen Produkten', 'Pflanzliche Alternativen'),
    ('Alkoholhaltige Getränke', 'Getränke mit Alkohol'),
    ('Alkoholfreie Getränke', 'Getränke'),
    ('Verschiedenes/Salz, Gewürze', 'Gewürze'),
    ('Verschiedenes/Saucen', 'Fette, Öle und Saucen'),
    ('Verschiedenes', 'Zutaten'),
]

# Gerichte, die in der Datenbank unter anderen Kategorien stehen (gehören in Stufe 4).
DISHES = {'Cordon bleu', 'Saltimbocca', 'Vitello tonnato', 'Ratatouille', 'Risotto', 'Gerstotto', 'Hirsotto',
          'Selleriescheiben', 'Kalbsgeschnetzeltes', 'Pouletgeschnetzeltes', 'Hackplätzli', 'Rindstatar'}

# Namen, die nach dem Zusammenfassen nicht gut lesbar oder missverständlich sind.
RENAME = {
    'Kaffee löslich': 'Instant-Kaffee',
    'Kochsalz': 'Salz',
    'Reis unpoliert': 'Vollkornreis',
    'Hühnerei': 'Ei',
    'Hühnereigelb': 'Eigelb',
    'Hühnereiweiss': 'Eiweiss',
    'Couscous-Körner': 'Couscous',
    'Kokosnussmilch': 'Kokosmilch',
    'Soja': 'Texturiertes Soja',
    'Sesamsamen': 'Sesam',
    'Rapsöl HOLL': 'Rapsöl',
    'Sonnenblumenöl HO': 'Sonnenblumenöl',
    'Palmöl - Palmfett': 'Palmöl',
}
# Pflanzliche Alternativen: „Burger, vegan …“ heisst „Veganer Burger“, nicht „Burger“.
VEGAN = {'Burger': 'Veganer Burger', 'Gehacktes': 'Veganes Gehacktes', 'Geschnetzeltes': 'Veganes Geschnetzeltes',
         'Nugget': 'Vegane Nuggets', 'Wurst': 'Vegane Wurst', 'Gehacktes/Geschnetzeltes/Pätzli': 'Mykoprotein'}
# Gruppe von Hand, wo die der Datenbank in der Suche verwirrt („Olive · Nüsse und Samen“)
GROUP_OVERRIDE = {'Olive': 'Gemüse', 'Avocado': 'Gemüse', 'Kokosmilch': 'Pflanzliche Alternativen'}
# Nachgestellte Eigenschaften, die für die Suche nichts unterscheiden („Reis poliert“ → Reis)
STRIP_WORDS = {'poliert', 'parboiled', 'unverdünnt', 'zubereitet', 'fettreduziert', 'gebraten', 'ungeschält', 'trocken',
               'weiss', 'rot', 'nera', 'dunkel', 'hell', 'französisch', 'italienisch', 'nature', 'roh', 'gekocht'}
# Umgangssprachliche Namen, die in der Datenbank fehlen
EXTRA_SYNONYMS = {
    'Mandelgetränk': ['Mandelmilch'],
    'Hafergetränk': ['Hafermilch'],
    'Sojagetränk': ['Sojamilch'],
    'Reisgetränk': ['Reismilch'],
    'Kokosnussgetränk': ['Kokosdrink'],
    'Wein': ['Rotwein', 'Weisswein'],
    'Schokolade': ['Schoggi'],
}
# Ergänzung Team (nicht in der Datenbank, aber oft genannt): Kräuter, Gewürze, „Alkohol“ allgemein.
# Sie tragen "team": true, damit klar bleibt, was nicht vom BLV stammt.
SUPPLEMENT = [
    ('Alkohol', 'Getränke mit Alkohol', ['Alkoholische Getränke']),
    ('Koriander', 'Kräuter', ['Cilantro']),
    ('Dill', 'Kräuter', []),
    ('Oregano', 'Kräuter', []),
    ('Majoran', 'Kräuter', []),
    ('Estragon', 'Kräuter', []),
    ('Liebstöckel', 'Kräuter', ['Maggikraut']),
    ('Pfeffer', 'Gewürze', []),
    ('Chili', 'Gewürze', ['Peperoncini', 'Peperoncino']),
    ('Curry', 'Gewürze', ['Currypulver']),
    ('Kreuzkümmel', 'Gewürze', ['Cumin']),
    ('Kümmel', 'Gewürze', []),
    ('Muskatnuss', 'Gewürze', ['Muskat']),
    ('Gewürznelken', 'Gewürze', ['Nelken']),
    ('Vanille', 'Gewürze', []),
    ('Safran', 'Gewürze', []),
    ('Knoblauchpulver', 'Gewürze', []),
]
STOPWORDS = {'Gewürz', 'Konserve', 'Durchschnitt', 'Schweiz', 'Neuseeland', 'Eidotter', 'Eiklar'}


def group_of(categories):
    for prefix, group in GROUPS:
        if any(category.startswith(prefix) for category in categories):
            return group
    return 'Zutaten'


def strip_parens(text):
    return re.sub(r'\([^)]*\)?', '', text)


def base_name(name):
    plain = strip_parens(name)
    parts = [part.strip() for part in plain.split(',')]
    first = re.split(r'\s+(?:in|im|mit|ohne|auf|aus|an)\s+', parts[0])[0]
    first = first.replace('/-', '').strip(' -')
    words = first.split()
    # ab der ersten Angabe mit Zahl („40 vol%“, „1.5% Fett“) abschneiden
    for i, word in enumerate(words):
        if re.search(r'\d', word):
            words = words[:i]
            break
    # nachgestellte kleine Eigenschaftswörter weglassen („Reis poliert“, „Sirup unverdünnt“), nicht „Mont d'Or“
    while len(words) > 1 and words[-1] in STRIP_WORDS:
        words = words[:-1]
    first = ' '.join(words)
    second = parts[1] if len(parts) > 1 else ''
    if first in VEGAN and second.startswith(('vegan', 'vegetarisch')):
        return VEGAN[first]
    return RENAME.get(first, first)


# Nur hier nennt der zweite Namensteil eine Sorte oder ein Stück („Rind, Entrecôte“, „Alge, Nori“);
# anderswo ist es eine Geschmacksrichtung („Cookie, Schoggi“, „Joghurt, Haselnuss“) und kein Suchbegriff.
CUT_GROUPS = ('Fleisch und Innereien', 'Fisch', 'Gemüse')


def extra_terms(name, categories):
    # „Rind, Entrecôte, roh“ → „Entrecôte“; „Spirituosen (z.B. Vodka, Gin)“ → „Vodka“, „Gin“:
    # Teile mit grossem Anfangsbuchstaben, ohne Zahlen
    cuts = any(category.startswith(CUT_GROUPS) for category in categories)
    parts = [strip_parens(part).strip() for part in name.split(',')[1:]] if cuts else []
    for inner in re.findall(r'\(([^)]*)\)?', name):
        parts += [part.replace('z.B.', '').strip() for part in inner.split(',')]
    cleaned = [re.sub(r'[()]', '', part).strip() for part in parts]
    return [part for part in cleaned
            if part and part[0].isupper() and not re.search(r'\d|%|/|"', part) and ' und ' not in part
            and len(part) <= 25 and part not in STOPWORDS]


def main(path):
    sheet = load_workbook(path, read_only=True)['Generische Lebensmittel']
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[2]
    col = {name: header.index(name) for name in ('Name', 'Synonyme', 'Kategorie')}
    version = rows[0][0]
    foods = {}
    for row in rows[3:]:
        name = (row[col['Name']] or '').strip()
        # Durchschnittswerte („Samen, Kerne, Nüsse (Durchschnitt)“) sind keine wählbaren Lebensmittel
        if not name or 'Durchschnitt' in name:
            continue
        categories = [part.strip() for part in str(row[col['Kategorie']] or '').split(';') if part.strip()]
        if any(category.startswith(EXCLUDED) for category in categories):
            continue
        base = base_name(name)
        if not base or ' und ' in base or len(base) < 2 or base in DISHES or base.split()[0] in {'Cordon', 'Vitello'}:
            continue
        synonyms = [s.strip() for s in str(row[col['Synonyme']] or '').split(';') if s.strip()]
        entry = foods.setdefault(base, {'name': base, 'category': GROUP_OVERRIDE.get(base, group_of(categories)), 'synonyms': []})
        for term in [*synonyms, *extra_terms(name, categories)]:
            term = RENAME.get(term, term)
            if term != base and term not in entry['synonyms'] and ',' not in term:
                entry['synonyms'].append(term)
    for base, extra in EXTRA_SYNONYMS.items():
        if base in foods:
            foods[base]['synonyms'] += [term for term in extra if term not in foods[base]['synonyms']]
    for name, group, synonyms in SUPPLEMENT:
        if name not in foods:
            foods[name] = {'name': name, 'category': group, 'synonyms': list(synonyms), 'team': True}
    items = sorted(foods.values(), key=lambda food: food['name'].lower())
    for item in items:
        if not item['synonyms']:
            del item['synonyms']
    doc = {
        'source': {
            'name': 'Schweizer Nährwertdatenbank, Generische Lebensmittel',
            'version': version,
            'publisher': 'Bundesamt für Lebensmittelsicherheit und Veterinärwesen (BLV)',
            'url': 'https://naehrwertdaten.ch',
            'license': 'Frei verwendbar mit Quellenvermerk',
            'note': 'Zutaten und Getränke, Varianten zusammengefasst (scripts/import-foods.py). Gerichte gehören in Stufe 4. Einträge mit "team": true sind eine Ergänzung des Teams (häufige Kräuter, Gewürze, Alkohol), nicht aus der Datenbank.',
        },
        'foods': items,
    }
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print(f'{len(items)} Lebensmittel: {OUT}')


if __name__ == '__main__':
    main(sys.argv[1])
